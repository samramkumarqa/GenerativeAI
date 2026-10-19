"""Read a BRD from a PDF, Word file, text file or web page into plain text."""

import io
import ipaddress
import re
import socket
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pypdf import PdfReader

from . import config


class DocumentError(ValueError):
    """Raised when a document can't be read or has too little text."""


def _normalise(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _check(text: str) -> str:
    text = _normalise(text)
    if len(text) < config.MIN_DOCUMENT_CHARS:
        raise DocumentError(
            "Not enough text found. If this is a scanned PDF, run it through OCR first."
        )
    return text[: config.MAX_DOCUMENT_CHARS]


def read_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # pypdf raises many types on malformed files
        raise DocumentError(f"Could not read PDF: {exc}") from exc
    return _check("\n\n".join(pages))


def read_docx(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError(f"Could not read Word file: {exc}") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _check("\n".join(parts))


def read_text(data: bytes) -> str:
    return _check(data.decode("utf-8", errors="replace"))


def read_upload(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        return read_pdf(data)
    if name.endswith(".docx"):
        return read_docx(data)
    if name.endswith((".txt", ".md")):
        return read_text(data)
    raise DocumentError("Unsupported file type. Use PDF, DOCX, TXT or MD.")


def _is_public_host(host: str) -> bool:
    """Block localhost / private ranges so the public demo can't be used to probe internal networks."""
    try:
        infos = socket.getaddrinfo(host, None)
    except socket.gaierror:
        return False
    for info in infos:
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
            return False
    return True


def read_url(url: str, timeout: int = 15) -> str:
    parsed = urlparse(url.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise DocumentError("Enter a full http(s):// URL.")
    if not _is_public_host(parsed.hostname):
        raise DocumentError("That address isn't a public website.")

    try:
        resp = requests.get(
            url,
            timeout=timeout,
            headers={"User-Agent": "Mozilla/5.0 (AI Test Case Generator)"},
        )
        resp.raise_for_status()
    except requests.RequestException as exc:
        raise DocumentError(f"Could not fetch URL: {exc}") from exc

    if "application/pdf" in resp.headers.get("Content-Type", "") or url.lower().endswith(".pdf"):
        return read_pdf(resp.content)

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form", "noscript"]):
        tag.decompose()
    root = soup.find("main") or soup.find("article") or soup.body or soup
    return _check(root.get_text(separator="\n"))
