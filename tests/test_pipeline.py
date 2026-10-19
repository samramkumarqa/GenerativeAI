import io

import pytest
from openpyxl import load_workbook

from fakes import FakeLLM, HashEmbedder
from tcgen import config
from tcgen.export import build_workbook
from tcgen.loaders import DocumentError, read_upload, read_url
from tcgen.pipeline import (
    _assign_ids,
    document_ids,
    extract_requirements,
    generate_test_cases,
    missing_ids,
    ready_to_export,
    review_summary,
    revise_test_cases,
    traceability,
)
from tcgen.rag import DocumentIndex
from tcgen.schemas import ExtractedRequirement, Requirement


def test_sample_pdf_has_requirements(sample_text):
    assert "FR-01" in sample_text and "NFR-04" in sample_text


def test_rejects_empty_and_unknown_files():
    with pytest.raises(DocumentError):
        read_upload("a.txt", b"too short")
    with pytest.raises(DocumentError):
        read_upload("a.exe", b"x" * 500)


@pytest.mark.parametrize("url", ["ftp://example.com/x", "http://localhost/brd", "http://127.0.0.1/x", "nonsense"])
def test_url_loader_blocks_non_public(url):
    with pytest.raises(DocumentError):
        read_url(url)


ALL_SAMPLE_IDS = [f"FR-0{i}" for i in range(1, 10)] + [f"NFR-0{i}" for i in range(1, 5)]


def test_document_ids_found(sample_text):
    assert document_ids(sample_text) == ALL_SAMPLE_IDS


def test_extraction_dedupes_and_recovers_missing(sample_text):
    llm = FakeLLM()
    reqs = extract_requirements(sample_text, llm)
    ids = [r.id for r in reqs]
    # First pass found FR-01 (reported twice) and one un-ID'd requirement;
    # the completeness check recovered every other ID written in the BRD.
    assert ids[:2] == ["FR-01", "REQ-001"]
    assert set(ids[2:]) == set(ALL_SAMPLE_IDS) - {"FR-01"}
    assert "ZZ-99" not in ids  # recovery can't add IDs that weren't asked for
    assert reqs[0].source_excerpt == "below INR 199"
    assert missing_ids(sample_text, reqs) == []
    assert sum("previous pass missed" in c[0] for c in llm.calls) == 2  # 12 missing IDs, 6 per call


def test_split_requirement_keeps_brd_id():
    def item(sid, title):
        return ExtractedRequirement(
            source_id=sid, title=title, description="d", category="Functional",
            priority="High", source_excerpt="",
        )

    reqs = _assign_ids([item("FR-03", "Slot required"), item("FR-03", "Slot window"),
                        item("FR-03", "Full slots greyed"), item("FR-04", "Fee")])
    assert [r.id for r in reqs] == ["FR-03.1", "FR-03.2", "FR-03.3", "FR-04"]
    assert missing_ids("FR-03 and FR-04 and FR-05", reqs) == ["FR-05"]


def test_retrieval_finds_relevant_chunk(sample_text):
    index = DocumentIndex(sample_text, HashEmbedder())
    context = index.context_for(["promo code discount alphanumeric"], k=1)
    assert "promo code" in context.lower()


def _reqs(n):
    return [
        Requirement(id=f"R{i}", title=f"Req {i}", description="d", category="Functional", priority="High")
        for i in range(1, n + 1)
    ]


def test_generation_batches_filters_and_numbers(sample_text):
    llm = FakeLLM()
    reqs = _reqs(5)
    tcs = generate_test_cases(reqs, DocumentIndex(sample_text, HashEmbedder()), llm, ["Functional"], 2)

    assert len(llm.calls) == -(-5 // config.GENERATION_BATCH_SIZE)
    assert [tc.requirement_id for tc in tcs] == ["R1", "R2", "R3", "R4", "R5"]  # hallucinated one dropped, order restored
    assert [tc.id for tc in tcs] == ["TC-001", "TC-002", "TC-003", "TC-004", "TC-005"]
    assert tcs[0].steps == ["Open cart", "Tap Checkout"]  # blank step removed
    assert "2 test cases for EACH requirement" in llm.calls[0][0]


def test_traceability_flags_uncovered(sample_text):
    reqs = _reqs(2)
    tcs = generate_test_cases(reqs[:1], DocumentIndex(sample_text, HashEmbedder()), FakeLLM(), ["Functional"])
    rows = traceability(reqs, tcs)
    assert [r["Coverage"] for r in rows] == ["Covered", "Not covered"]
    assert rows[0]["Test Case IDs"] == "TC-001"


def _review(tcs, statuses, comments=None):
    comments = comments or {}
    return [
        tc.model_copy(update={"review_status": st, "reviewer_comment": comments.get(tc.id, "")})
        for tc, st in zip(tcs, statuses)
    ]


def test_workbook_exports_only_approved_with_audit(sample_text):
    reqs = _reqs(3)
    tcs = generate_test_cases(reqs, DocumentIndex(sample_text, HashEmbedder()), FakeLLM(), ["Functional"])
    tcs = _review(tcs, ["Approved", "Rejected", "Approved"], {"TC-002": "Too vague"})
    wb = load_workbook(io.BytesIO(build_workbook(reqs, tcs, "sample.pdf", "fake", reviewer="Ram")))

    assert wb.sheetnames == ["Summary", "Test Cases", "Traceability", "Requirements", "Rejected"]
    cases = wb["Test Cases"]
    assert [cases.cell(row=r, column=1).value for r in range(2, cases.max_row + 1)] == ["TC-001", "TC-003"]
    assert cases["G2"].value == "1. Open cart\n2. Tap Checkout"
    assert cases["J2"].value == "Not Run"
    assert cases.data_validations.dataValidation  # execution status dropdown present
    assert wb["Rejected"]["A2"].value == "TC-002" and wb["Rejected"]["G2"].value == "Too vague"
    assert wb["Traceability"]["H3"].value == "Not covered"  # R2's only test was rejected

    summary = {wb["Summary"].cell(row=r, column=1).value: wb["Summary"].cell(row=r, column=2).value
               for r in range(3, 20)}
    assert summary["Reviewed and approved by"] == "Ram"
    assert summary["Rejected by reviewer"] == 1
    assert summary["Coverage"] == "=IF(B7=0,0,B10/B7)"  # Requirements row 7, covered row 10


def test_export_gate():
    tcs = _review(generate_test_cases(_reqs(2), _index(), FakeLLM(), ["Functional"]), ["Approved", "Pending"])
    assert ready_to_export(tcs, "Ram") == (False, "1 test case(s) still pending review.")
    tcs = _review(tcs, ["Rejected", "Rejected"])
    assert ready_to_export(tcs, "Ram") == (False, "Approve at least one test case.")
    tcs = _review(tcs, ["Approved", "Rejected"])
    assert ready_to_export(tcs, "  ") == (False, "Enter the reviewer's name to sign off.")
    assert ready_to_export(tcs, "Ram") == (True, "")
    assert review_summary(tcs) == {"Pending": 0, "Approved": 1, "Rejected": 1}


def test_revise_rejected_uses_feedback():
    reqs = _reqs(2)
    llm = FakeLLM()
    tcs = generate_test_cases(reqs, _index(), llm, ["Functional"])
    tcs = _review(tcs, ["Rejected", "Rejected"], {"TC-001": "Add an invalid amount case"})

    out = revise_test_cases(tcs, reqs, _index(), llm, next_number=3)

    assert "REVIEWER FEEDBACK: Add an invalid amount case" in llm.calls[-1][1]
    assert "REVIEWER FEEDBACK: (none given)" in llm.calls[-1][1]
    new = out["TC-001"]
    assert (new.id, new.requirement_id, new.revision, new.review_status) == ("TC-003", "R1", 2, "Pending")
    assert "Add an invalid amount case" in new.title
    assert out["TC-002"].id == "TC-004"


def _index():
    from conftest import SAMPLE_BRD
    from tcgen.loaders import read_upload

    return DocumentIndex(read_upload(SAMPLE_BRD.name, SAMPLE_BRD.read_bytes()), HashEmbedder())
